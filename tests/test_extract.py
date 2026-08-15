"""Node `extract`. Phần suy luận cấu trúc là THUẦN PYTHON nên test được thật.

Chỉ lượt ánh xạ header mới gọi LLM, và nó bị thay bằng stub ở đây — không test
nào trong file này đi ra mạng.
"""

import asyncio
import uuid
from datetime import date, datetime
from pathlib import Path

import pytest
from langchain_core.documents import Document
from openpyxl import Workbook

from TAR_agent.graph_admin.helpers.loaders import read_workbook
from TAR_agent.graph_admin.helpers.rows import HeaderMap
from TAR_agent.graph_admin.nodes.extract import (
    Extract,
    _ChunkIndex,
    map_columns,
    split_rows,
)
from TAR_agent.graph_admin.state import UploadedFile

# Ánh xạ cột dùng chung cho phần lớn test: đúng thứ tự cột của `lich_file`.
COLS = {
    "tt": 0, "cong_viec": 1, "don_vi": 2, "ngay_bd": 3,
    "ngay_ht": 4, "so_ngay": 5, "ghi_chu": 6,
}

HEADER = ["TT", "Hạng mục thực hiện", "Đơn vị thực hiện", "Ngày bắt đầu",
          "Ngày hoàn thành", "Số ngày", "Ghi chú"]


@pytest.fixture
def lich_file(tmp_path: Path) -> Path:
    """Bảng lịch công việc phân cấp, đúng hình dạng file thật.

    Có đủ: dòng tiêu đề báo cáo chiếm một ô ở trên, mã cấp I / I.1 / I.2 /
    I.2.1, dòng "-" mang ngày, và một dòng trống ở giữa.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lịch"
    sheet.append(["BẢNG TIẾN ĐỘ DỰ ÁN X"])
    sheet.append(HEADER)
    sheet.append(["I", "Chuẩn bị đầu tư", None, None, None, None, None])
    sheet.append(["I.1", "Lập chủ trương", None, None, None, None, None])
    sheet.append(["-", "Khảo sát hiện trạng", "Ban QLDA",
                  date(2026, 6, 1), date(2026, 6, 5), 5, "Đã xong"])
    sheet.append([None] * 7)
    sheet.append(["I.2", "Thẩm định", None, None, None, None, None])
    sheet.append(["I.2.1", "Thẩm định kỹ thuật", None, None, None, None, None])
    sheet.append(["-", "Lấy ý kiến sở ngành", "Sở KHĐT",
                  date(2026, 6, 6), datetime(2026, 6, 20), None, None])
    sheet.append(["II", "Thực hiện đầu tư", None, None, None, None, None])
    sheet.append(["-", "Thi công", "Nhà thầu", date(2026, 7, 1), None, None, None])

    path = tmp_path / "TiendoT6.xlsx"
    workbook.save(path)
    workbook.close()
    return path


class TestSplitRows:
    """Phân loại tiêu đề / dòng thật bằng cột TT. Không LLM — mã TT là ký hiệu
    cấu trúc, không cần suy luận ngôn ngữ."""

    def test_dong_MA_CAP_khong_sinh_dong_cong_viec(self, lich_file: Path):
        rows = split_rows(read_workbook(lich_file)[0].rows, COLS)
        assert [r["cong_viec"] for r in rows] == [
            "Khảo sát hiện trạng", "Lấy ý kiến sở ngành", "Thi công",
        ]

    def test_dong_that_mang_ngu_canh_GAN_NHAT_phia_tren(self, lich_file: Path):
        rows = split_rows(read_workbook(lich_file)[0].rows, COLS)
        assert (rows[0]["giai_doan"], rows[0]["nhom"], rows[0]["nhom_con"]) == (
            "Chuẩn bị đầu tư", "Lập chủ trương", None,
        )
        assert (rows[1]["giai_doan"], rows[1]["nhom"], rows[1]["nhom_con"]) == (
            "Chuẩn bị đầu tư", "Thẩm định", "Thẩm định kỹ thuật",
        )

    def test_ma_cap_NONG_HON_reset_cac_cap_con(self, lich_file: Path):
        """"II" là cấp 1, phải xoá "Thẩm định" / "Thẩm định kỹ thuật" của I.2.1.

        Thiếu bước reset thì "Thi công" mang theo nhóm con của một giai đoạn
        khác hẳn — sai mà không có triệu chứng nào ngoài câu trả lời lạ.
        """
        rows = split_rows(read_workbook(lich_file)[0].rows, COLS)
        thi_cong = rows[-1]
        assert thi_cong["giai_doan"] == "Thực hiện đầu tư"
        assert thi_cong["nhom"] is None
        assert thi_cong["nhom_con"] is None

    def test_doc_dung_du_lieu_rieng_cua_dong(self, lich_file: Path):
        rows = split_rows(read_workbook(lich_file)[0].rows, COLS)
        assert rows[0]["don_vi"] == "Ban QLDA"
        assert rows[0]["ngay_bd"] == date(2026, 6, 1)
        assert rows[0]["so_ngay"] == 5
        assert rows[0]["ghi_chu"] == "Đã xong"

    def test_datetime_cua_excel_thanh_date(self, lich_file: Path):
        rows = split_rows(read_workbook(lich_file)[0].rows, COLS)
        assert rows[1]["ngay_ht"] == date(2026, 6, 20)

    def test_dong_co_MA_CAP_nhung_CO_NGAY_van_la_dong_that(self):
        """File thật hay đánh số cả dòng công việc. Có ngày thì nó mang dữ liệu."""
        rows = split_rows(
            [["1.1", "Khảo sát", "Ban QLDA", date(2026, 6, 1), None, None, None]], COLS
        )
        assert [r["cong_viec"] for r in rows] == ["Khảo sát"]

    def test_thieu_cot_cong_viec_thi_tra_RONG(self):
        """Thiếu đúng cột đó thì mọi dòng đều trượt NOT NULL — đọc tiếp để vứt đi."""
        assert split_rows([["-", "Khảo sát"]], {"tt": 0}) == []

    def test_dong_ngan_hon_so_cot_khong_lam_no_IndexError(self):
        """openpyxl không đệm dòng cho đủ bề rộng sheet."""
        rows = split_rows([["-", "Khảo sát"]], COLS)
        assert rows[0]["cong_viec"] == "Khảo sát"
        assert rows[0]["ghi_chu"] is None

    def test_o_rong_thanh_None_khong_thanh_chuoi_rong(self):
        """`ket_qua_dau_ra IS NULL` là một trong những câu query_data phải trả
        lời được — `''` và `NULL` khác nhau ở SQL."""
        rows = split_rows([["-", "Khảo sát", "   ", None, None, None, ""]], COLS)
        assert rows[0]["don_vi"] is None
        assert rows[0]["ghi_chu"] is None


class TestMapColumns:
    def test_khop_nguyen_van(self):
        got = map_columns(HEADER, HeaderMap(tt="TT", cong_viec="Hạng mục thực hiện"))
        assert got == {"tt": 0, "cong_viec": 1}

    def test_khop_lech_hoa_thuong_va_mat_dau(self):
        """Model chép tên cột "gần đúng" là chuyện thường."""
        got = map_columns(HEADER, HeaderMap(cong_viec="hang muc thuc hien"))
        assert got == {"cong_viec": 1}

    def test_cot_khong_do_ra_thi_VANG_MAT_chu_khong_gan_bua(self):
        got = map_columns(HEADER, HeaderMap(cong_viec="Cột không tồn tại"))
        assert got == {}

    def test_null_bi_bo_qua(self):
        assert map_columns(HEADER, HeaderMap()) == {}


class TestChunkIndex:
    def test_tim_dung_chunk_chua_ten_cong_viec(self):
        doc_id = uuid.uuid4()
        index = _ChunkIndex(
            doc_id,
            [Document(page_content="phần đầu"), Document(page_content="Khảo sát hiện trạng")],
        )
        from TAR_agent.utils.text import chunk_uuid

        assert index.locate("Khảo sát hiện trạng") == chunk_uuid(doc_id, 1)

    def test_khong_tim_thay_thi_None_chu_KHONG_gan_bua(self):
        """chunk_id dùng để truy ngược số về nguồn — gán bừa tệ hơn để trống."""
        index = _ChunkIndex(uuid.uuid4(), [Document(page_content="phần đầu")])
        assert index.locate("Việc không có ở đâu") is None
        assert index.locate(None) is None
        assert index.locate("") is None


class _FakeMapper:
    """Thay client LLM ánh xạ header. Trả sẵn một HeaderMap, hoặc ném."""

    def __init__(self, result: HeaderMap | None, boom: bool = False) -> None:
        self.result, self.boom, self.calls = result, boom, 0

    async def ainvoke(self, _payload):
        self.calls += 1
        if self.boom:
            raise RuntimeError("Gemini 429")
        return self.result


def _state(path: Path) -> dict:
    return {
        "upload": UploadedFile("TiendoT6.xlsx", path, path.stat().st_size),
        "project_id": uuid.uuid4(),
        "base_name": "TiendoT6.xlsx",
        "file_kind": "xlsx",
        "as_of_date": date(2026, 6, 30),
        "chunks": [Document(page_content="Khảo sát hiện trạng | Ban QLDA")],
    }


FULL_MAP = HeaderMap(
    tt="TT", cong_viec="Hạng mục thực hiện", don_vi="Đơn vị thực hiện",
    ngay_bd="Ngày bắt đầu", ngay_ht="Ngày hoàn thành", so_ngay="Số ngày",
    ghi_chu="Ghi chú",
)


class TestNodeXlsx:
    def test_MOT_luot_LLM_moi_sheet_du_bang_bao_nhieu_dong(self, lich_file: Path):
        """Một file 500 dòng tốn 1 lượt chứ không phải 500 — và số liệu KHÔNG
        đi qua model nên không có đường nào để sai."""
        node = Extract()
        mapper = _FakeMapper(FULL_MAP)
        node.header_mapper = mapper

        out = asyncio.run(node(_state(lich_file)))
        assert mapper.calls == 1
        assert [r.cong_viec for r in out["rows"]] == [
            "Khảo sát hiện trạng", "Lấy ý kiến sở ngành", "Thi công",
        ]
        assert out["extract_error"] is None

    def test_so_ngay_TINH_bang_python(self, lich_file: Path):
        node = Extract()
        node.header_mapper = _FakeMapper(FULL_MAP)
        rows = asyncio.run(node(_state(lich_file)))["rows"]
        assert rows[0].so_ngay == 5
        assert rows[1].so_ngay == 15  # 06/6 -> 20/6, tính cả ngày đầu

    def test_bao_lai_anh_xa_cot_da_dung_de_admin_soat(self, lich_file: Path):
        """Header lạ làm ánh xạ hỏng cho CẢ sheet, không phải vài dòng — in ra
        là cách duy nhất admin phát hiện."""
        node = Extract()
        node.header_mapper = _FakeMapper(FULL_MAP)
        out = asyncio.run(node(_state(lich_file)))
        assert "cong_viec=Hạng mục thực hiện" in out["column_map"]["Lịch"]

    def test_LLM_hong_thi_KHONG_tra_error_va_KHONG_huy_luot_nap(self, lich_file: Path):
        """`extract` là node duy nhất của nhánh nạp không có cạnh về `report`.
        Trả `error` ở đây là huỷ luôn phần chunk đang chạy tốt."""
        node = Extract()
        node.header_mapper = _FakeMapper(None, boom=True)

        out = asyncio.run(node(_state(lich_file)))
        assert "error" not in out
        assert out["rows"] == []
        assert out["extract_error"] is None, "lỗi một sheet không phải lỗi cả lượt"
        assert "không ánh xạ được cột" in out["rows_rejected"][0]

    def test_sheet_khong_phai_bang_lich_thi_bo_qua_lang_le(self, lich_file: Path):
        """Sheet vướng mắc / ký nhận: model trả toàn null theo luật 4 của prompt."""
        node = Extract()
        node.header_mapper = _FakeMapper(HeaderMap())
        out = asyncio.run(node(_state(lich_file)))
        assert out["rows"] == []
        assert out["rows_rejected"] == [], "không phải lỗi thì đừng làm ồn"

    def test_dem_dong_bi_loai_bang_DO_LECH(self, lich_file: Path):
        node = Extract()
        node.header_mapper = _FakeMapper(FULL_MAP)
        out = asyncio.run(node(_state(lich_file)))
        assert out["rows_rejected_count"] == 0

    def test_chunk_id_gan_dung_doan_sinh_ra_dong(self, lich_file: Path):
        node = Extract()
        node.header_mapper = _FakeMapper(FULL_MAP)
        rows = asyncio.run(node(_state(lich_file)))["rows"]
        # Chỉ chunk duy nhất trong state chứa "Khảo sát hiện trạng".
        assert rows[0].chunk_id is not None
        assert rows[2].chunk_id is None, "không truy được thì để trống"


class TestNodeKhongCoGiDeLam:
    def test_khong_co_chunk_thi_tra_rong_khong_no(self, lich_file: Path):
        node = Extract()
        state = _state(lich_file) | {"chunks": []}
        out = asyncio.run(node(state))
        assert out["rows"] == [] and out["extract_error"] is None
