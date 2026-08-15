"""File mẫu và kho GIẢ cho test.

File mẫu dựng bằng code chứ không để sẵn trong repo: một file .xlsx commit vào
git là 10KB nhị phân không ai diff được, và khi test hỏng thì không biết nội
dung nó ra sao. Dựng bằng code thì đọc test là thấy đúng dữ liệu đang được kiểm.

Kho giả (`FakeKho`) thay toàn bộ tầng DB và embedding: test luồng admin chạy
KHÔNG cần Neon, không cần GOOGLE_API_KEY, không tốn tiền embedding và không phụ
thuộc mạng. Đổi lại nó không bắt được lỗi SQL — phần đó phải kiểm bằng tay trên
DB thật.
"""

import asyncio
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook



@pytest.fixture
def txt_file(tmp_path: Path) -> Path:
    path = tmp_path / "GhiChu.txt"
    path.write_text(
        "# Vướng mắc tháng 6\n\n"
        "Hạng mục 3 chậm do chờ mặt bằng.\n"
        "Đội xe đã bố trí đủ phương tiện.\n",
        encoding="utf-8",
    )
    return path


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    """Bảng tiến độ có đủ những kiểu ô hay làm hỏng chuyện.

    Dòng tiêu đề chung ở trên, ngày tháng, phần trăm nguyên, ô trống, ô có ký tự
    `|`, và một dòng hoàn toàn trống ở giữa.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hạng mục"
    sheet.append(["BÁO CÁO TIẾN ĐỘ THÁNG 6"])
    sheet.append(["Mã", "Tên hạng mục", "Bắt đầu", "Kết thúc KH", "% hoàn thành"])
    sheet.append(["HM1", "Khảo sát", date(2026, 6, 1), datetime(2026, 6, 15), 100])
    sheet.append(["HM2", "Thiết kế | bản vẽ", date(2026, 6, 10), None, 40])
    sheet.append([None, None, None, None, None])
    sheet.append(["HM3", "Thi công", date(2026, 6, 20), datetime(2026, 9, 30), 0])

    second = workbook.create_sheet("Vướng mắc")
    second.append(["Nội dung", "Đơn vị"])
    second.append(["Chờ mặt bằng", "Ban QLDA"])

    path = tmp_path / "TiendoT6.xlsx"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def empty_xlsx_file(tmp_path: Path) -> Path:
    workbook = Workbook()
    path = tmp_path / "Rong.xlsx"
    workbook.save(path)
    workbook.close()
    return path




@dataclass
class FakeDocument:
    """Đủ những thuộc tính mà check_file và report đọc tới."""

    document_id: uuid.UUID
    project_id: uuid.UUID
    file_name: str
    content_sha256: str
    chunk_count: int
    as_of_date: date = date(2026, 6, 30)
    uploaded_at: datetime = field(
        default_factory=lambda: datetime(2026, 7, 5, tzinfo=timezone.utc)
    )


@dataclass
class StoredCall:
    """Một lời gọi writer.save đã ghi nhận, để test soi lại."""

    project_id: uuid.UUID
    file_name: str
    file_kind: str
    chunk_count: int
    as_of_date: date
    uploaded_by: int
    row_count: int = 0


class FakeKho:
    """Thay tầng DB + embedding. Đếm luôn số lần gọi để test khẳng định
    "KHÔNG embed khi admin bấm Huỷ" — thứ tốn tiền thật nếu làm sai.
    """

    def __init__(self):
        self.projects: dict[uuid.UUID, tuple[str, int]] = {}
        self.documents: dict[tuple[uuid.UUID, str], FakeDocument] = {}
        self.embed_calls = 0
        self.extract_calls = 0
        # Thứ node `extract` giả sẽ trả về. Test nào cần kiểm phần ghi dòng thì
        # gán vào đây trước khi chạy lượt nạp.
        self.extract_rows: list = []
        self.save_calls: list[StoredCall] = []

    # --- dựng dữ liệu cho test ---

    def add_project(self, name: str, document_count: int = 0) -> uuid.UUID:
        from TAR_agent.utils.text import name_uuid

        pid = name_uuid(name)
        self.projects[pid] = (name, document_count)
        return pid

    def add_document(
        self, project_id: uuid.UUID, file_name: str, sha: str, chunk_count: int = 5
    ) -> FakeDocument:
        from TAR_agent.utils.text import document_uuid

        doc = FakeDocument(
            document_id=document_uuid(project_id, file_name),
            project_id=project_id,
            file_name=file_name,
            content_sha256=sha,
            chunk_count=chunk_count,
        )
        self.documents[(project_id, file_name)] = doc
        return doc

    # --- bản thay cho hàm thật ---

    async def list_projects(self):
        from TAR_agent.utils.projects import ProjectBrief

        return [
            ProjectBrief(pid, name, count)
            for pid, (name, count) in sorted(
                self.projects.items(), key=lambda kv: kv[1][0]
            )
        ]

    async def create_project(self, name: str):
        from TAR_agent.utils.projects import ProjectBrief
        from TAR_agent.utils.text import name_uuid, normalize_name

        if not normalize_name(name):
            raise ValueError(f"Tên rỗng sau chuẩn hoá: {name!r}")
        pid = name_uuid(name)
        if pid in self.projects:
            return ProjectBrief(pid, self.projects[pid][0], 0), False
        self.projects[pid] = (name, 0)
        return ProjectBrief(pid, name, 0), True

    def find_by_name(self, project_id: uuid.UUID, file_name: str):
        return self.documents.get((project_id, file_name))

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.embed_calls += 1
        return [[0.1] * 768 for _ in texts]

    async def extract(self, _state) -> dict:
        """Thay node `extract`. KHÔNG gọi Gemini.

        Đếm lượt gọi cùng lý do như `embed_calls`: `extract` chạy sau cả hai
        điểm dừng, nên "admin bấm Huỷ mà vẫn tốn một lượt LLM" là đúng loại lỗi
        không có triệu chứng nào ngoài hoá đơn cuối tháng.
        """
        self.extract_calls += 1
        return {
            "rows": list(self.extract_rows),
            "rows_rejected": [],
            "rows_rejected_count": 0,
            "column_map": {},
            "extract_error": None,
        }

    def save(self, **kwargs):
        from TAR_agent.graph_admin.helpers.writer import StoredDocument
        from TAR_agent.utils.text import document_uuid

        chunks = kwargs["chunks"]
        key = (kwargs["project_id"], kwargs["file_name"])
        old = self.documents.get(key)
        replaced = old.chunk_count if old else 0

        rows = kwargs.get("rows") or []
        self.save_calls.append(
            StoredCall(
                project_id=kwargs["project_id"],
                file_name=kwargs["file_name"],
                file_kind=kwargs["file_kind"],
                chunk_count=len(chunks),
                as_of_date=kwargs["as_of_date"],
                uploaded_by=kwargs["uploaded_by"],
                row_count=len(rows),
            )
        )
        doc_id = document_uuid(kwargs["project_id"], kwargs["file_name"])
        self.documents[key] = FakeDocument(
            document_id=doc_id,
            project_id=kwargs["project_id"],
            file_name=kwargs["file_name"],
            content_sha256=kwargs["content_sha256"],
            chunk_count=len(chunks),
            as_of_date=kwargs["as_of_date"],
        )
        return StoredDocument(doc_id, len(chunks), replaced, len(rows), 0)


@pytest.fixture
def kho(monkeypatch) -> FakeKho:
    """Thay mọi đường đi ra ngoài của graph admin bằng kho giả.

    Vá tại NƠI DÙNG chứ không tại nơi định nghĩa: các node dùng
    `from ... import ten_ham` nên tên đã được gắn vào namespace của node, vá
    module gốc không có tác dụng.

    Lấy module bằng import_module chứ không `from graph_admin.nodes import
    ask_project`: nodes/__init__.py xuất lại HÀM trùng tên với MODULE, nên cách
    viết kia trả về hàm và monkeypatch.setattr sẽ nổ AttributeError.
    """
    from importlib import import_module

    ask_project_node = import_module("TAR_agent.graph_admin.nodes.ask_project")
    handle_text_node = import_module("TAR_agent.graph_admin.nodes.handle_text")
    store_node = import_module("TAR_agent.graph_admin.nodes.store")
    extract_node = import_module("TAR_agent.graph_admin.nodes.extract")
    writer = import_module("TAR_agent.graph_admin.helpers.writer")
    doc_proc = import_module("persistence.proc.documents")

    fake = FakeKho()
    monkeypatch.setattr(ask_project_node, "list_projects", fake.list_projects)
    monkeypatch.setattr(handle_text_node, "list_projects", fake.list_projects)
    monkeypatch.setattr(handle_text_node, "create_project", fake.create_project)
    monkeypatch.setattr(doc_proc, "find_by_name", fake.find_by_name)
    monkeypatch.setattr(store_node, "embed_documents", fake.embed_documents)
    monkeypatch.setattr(writer, "save", fake.save)
    # Vá trên LỚP, không trên instance: `Extract()` được dựng lúc build graph
    # (fixture `graph`), mà thứ tự hai fixture phụ thuộc thứ tự tham số của
    # từng test — vá lớp thì đúng dù chúng chạy theo thứ tự nào.
    monkeypatch.setattr(
        extract_node.Extract, "__call__", lambda _self, state: fake.extract(state)
    )
    return fake


@pytest.fixture
def graph():
    """Graph admin với checkpointer trong RAM — không đụng Postgres."""
    from langgraph.checkpoint.memory import MemorySaver

    from TAR_agent.graph_admin.graph import build_admin_graph

    return build_admin_graph(MemorySaver())




@pytest.fixture
def upload_factory(tmp_path: Path):
    """Dựng UploadedFile y như app/telegram/download.py làm.

    File tạm mang tên NGẪU NHIÊN KHÔNG ĐUÔI — đúng điều kiện thật, để loader nào
    ngầm giả định "đường dẫn có đuôi .txt" sẽ lộ ra ở test chứ không phải lúc
    admin gửi file.
    """
    from TAR_agent.graph_admin.state import UploadedFile

    def make(file_name: str, content: str = "Nội dung báo cáo tiến độ.") -> UploadedFile:
        path = tmp_path / f"tg-{uuid.uuid4().hex}"
        path.write_text(content, encoding="utf-8")
        return UploadedFile(file_name, path, path.stat().st_size)

    return make


ADMIN_ID = 7_000_000_000


def turn(graph, cfg: dict, text: str = "", upload=None):
    """Một LƯỢT MỚI: đúng những gì handle_admin_message đẩy vào graph."""
    from langchain_core.messages import HumanMessage

    return asyncio.run(
        graph.ainvoke(
            {
                "messages": [HumanMessage(text)],
                "upload": upload,
                "uploaded_by": ADMIN_ID,
            },
            cfg,
        )
    )


def resume(graph, cfg: dict, value: dict):
    """Admin bấm nút: đúng những gì handle_admin_callback đẩy vào graph."""
    from langgraph.types import Command

    return asyncio.run(graph.ainvoke(Command(resume=value), cfg))


def waiting_for(graph, cfg: dict) -> dict | None:
    from app.routers.webhooks import pending_interrupt

    return asyncio.run(pending_interrupt(graph, cfg))


def kinds(out: dict) -> list[str]:
    return [e["kind"] for e in out.get("outbox", [])]
